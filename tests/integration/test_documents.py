"""Documents other than PDF (Word, PowerPoint, Excel, HTML, EPUB, Markdown, text): read with MarkItDown into the
same canonical document as a PDF, so search, export and the AI features work on them unchanged."""

from __future__ import annotations

import sqlite3
import zipfile
from pathlib import Path

import pytest

from litforge.config import Config, Paths
from litforge.export.markdown import export_markdown
from litforge.parsing.canonical import read_canonical
from litforge.parsing.markitdown_parser import document_from_markdown, is_document
from litforge.pipeline import add_pdfs
from litforge.projects import create_project
from litforge.ui import api

pytest.importorskip("markitdown")

MD = """# Coupling in whole-brain models

## Abstract

We tune the global coupling gain.

## Results

The gain was 3.25 across 40 subjects.

| Gain | Fit |
|---|---|
| 3.25 | 0.61 |
"""


def _docx(path: Path, paragraphs: list[tuple[str, str]]) -> Path:
    """A minimal Word file: (style, text) paragraphs, style 'Title', 'Heading1' or ''."""
    body = "".join(
        f'<w:p>{f"<w:pPr><w:pStyle w:val=\"{style}\"/></w:pPr>" if style else ""}<w:r><w:t>{text}</w:t></w:r></w:p>'
        for style, text in paragraphs)
    ns = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("[Content_Types].xml",
                   '<?xml version="1.0"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
                   '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
                   '<Default Extension="xml" ContentType="application/xml"/><Override PartName="/word/document.xml" '
                   'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
                   '</Types>')
        z.writestr("_rels/.rels",
                   '<?xml version="1.0"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/'
                   'relationships"><Relationship Id="r1" Type="http://schemas.openxmlformats.org/officeDocument/'
                   '2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>')
        z.writestr("word/document.xml", f'<?xml version="1.0"?><w:document {ns}><w:body>{body}</w:body></w:document>')
    return path


def test_markdown_becomes_sections_with_a_title() -> None:
    doc, title = document_from_markdown(MD, "markitdown", "0")
    assert title == "Coupling in whole-brain models"
    assert [(s.heading, s.kind) for s in doc.sections] == [
        ("Coupling in whole-brain models", "front"), ("Abstract", "abstract"), ("Results", "body")]
    results = doc.sections[2].paragraphs
    assert results[0].text == "The gain was 3.25 across 40 subjects."
    assert results[1].text.startswith("| Gain | Fit |")  # a table stays one block
    assert doc.page_count == 1 and all(p.page == 1 for s in doc.sections for p in s.paragraphs)


def test_document_types() -> None:
    assert all(is_document(Path(f"a{s}")) for s in (".docx", ".pptx", ".xlsx", ".html", ".epub", ".md", ".txt"))
    assert not is_document(Path("a.pdf")) and not is_document(Path("a.exe"))


def test_documents_are_added_read_exported_and_not_served_as_pdf(conn: sqlite3.Connection, paths: Paths,
                                                                 config: Config, tmp_path: Path) -> None:
    pid = create_project(conn, paths, "docs")
    folder = tmp_path / "mine"
    folder.mkdir()
    (folder / "notes.md").write_text(MD, encoding="utf-8")
    _docx(folder / "draft.docx", [("Title", "Neural mass models"), ("Heading1", "Methods"),
                                  ("", "We fitted the model to the recordings.")])
    (folder / "page.html").write_text("<html><head><title>A web page</title></head><body><h1>A web page</h1>"
                                      "<p>Text from the web.</p></body></html>", encoding="utf-8")
    (folder / "program.exe").write_bytes(b"MZ")  # not a document: ignored

    # A folder brings its Word file; Markdown and web pages are taken when chosen one by one.
    report, results, parsed = add_pdfs(conn, paths, config, pid, [folder, folder / "notes.md", folder / "page.html"],
                                       "my documents")
    assert sorted(Path(r.path).name for r in results) == ["draft.docx", "notes.md", "page.html"]
    assert {p.status for p in parsed} == {"ok"} and {p.parser for p in parsed} == {"markitdown"}
    titles = sorted(r["title"] for r in conn.execute(
        "SELECT p.title FROM project_papers pp JOIN papers p ON p.id = pp.paper_id WHERE pp.project_id=?", (pid,)))
    assert titles == ["A web page", "Coupling in whole-brain models", "Neural mass models"]

    row = conn.execute("""SELECT a.paper_id, a.pdf_path, r.output_dir FROM artifacts a
                          JOIN parse_runs r ON r.artifact_id = a.id AND r.is_current = 1
                          JOIN papers p ON p.id = a.paper_id WHERE p.title = 'Neural mass models'""").fetchone()
    assert row["pdf_path"].endswith(".docx")  # kept with its own extension
    assert "We fitted the model" in read_canonical(paths.root / row["output_dir"]).full_text()
    assert api._pdf_path(conn, paths, row["paper_id"]) is None  # never served or shown as a PDF

    out, count = export_markdown(conn, paths, "docs", pid, config=config)
    assert count == 3
    notes = "".join(p.read_text(encoding="utf-8") for p in (out / "knowledge_base").glob("*.md"))
    assert "We fitted the model to the recordings." in notes and "Text from the web." in notes
    assert "None" not in notes and "None" not in (out / "README.md").read_text()  # an unknown version is left out
    assert not list((out.parent / "pdfs").glob("*.pdf"))  # no Word file renamed to .pdf
    assert len(list((out.parent / "pdfs").glob("*.docx"))) == 1

    again, results2, _ = add_pdfs(conn, paths, config, pid, [folder / "notes.md"], "again")
    assert [r.status for r in results2] == ["duplicate"]


def _sample(folder: Path, name: str) -> Path:
    from openpyxl import Workbook
    from pptx import Presentation
    path = folder / name
    if name.endswith(".pptx"):
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[0])
        slide.shapes.title.text = "Deck title"
        slide.placeholders[1].text = "Results of the fit"
        deck.save(path)
    elif name.endswith(".xlsx"):
        book = Workbook()
        book.active.append(["Gain", "Fit"])
        book.active.append([3.25, 0.61])
        book.save(path)
    else:
        path.write_text("<html><head><title>Real Title</title></head><body><nav>Home | About</nav><h1>Heading</h1>"
                        "<p>Body <img src='https://tracker.example/x.png' alt='a chart'></p></body></html>")
    return path


def test_titles_of_slides_sheets_and_web_pages(tmp_path: Path) -> None:
    # Review: a deck's title was "<!-- Slide number: 1 -->", a workbook's "Sheet", a page's its menu text.
    from litforge.parsing.markitdown_parser import document_title
    assert document_title(_sample(tmp_path, "deck.pptx")) == "Deck title"
    assert document_title(_sample(tmp_path, "results-2024.xlsx")) is None  # the file name is used instead
    assert document_title(_sample(tmp_path, "page.html")) == "Real Title"


def test_remote_images_are_not_kept(tmp_path: Path) -> None:
    from litforge.parsing.markitdown_parser import markdown_of
    text = markdown_of(_sample(tmp_path, "page.html"))
    assert "tracker.example" not in text and "a chart" in text


def test_folders_bring_pdfs_and_office_documents_not_every_text_file(tmp_path: Path) -> None:
    # Review: a Zotero storage folder or an export folder turned every .html/.md/.txt into a paper.
    from litforge.ingest import iter_pdfs
    folder = tmp_path / "papers"
    (folder / ".hidden").mkdir(parents=True)
    for name in ("a.pdf", "b.docx", "c.pptx", "d.epub", "notes.txt", "README.md", "snapshot.html", ".hidden/e.docx"):
        (folder / name).write_bytes(b"x")
    assert sorted(p.name for p in iter_pdfs([folder], documents=True)) == ["a.pdf", "b.docx", "c.pptx", "d.epub"]
    loose = folder / "notes.txt"
    assert iter_pdfs([loose], documents=True) == [loose.resolve()]  # a file chosen by hand is always taken


def test_a_quarantined_file_can_be_added_again_as_a_document(conn: sqlite3.Connection, paths: Paths,
                                                             config: Config, tmp_path: Path) -> None:
    # Review: a Word file whose long name lost its extension was quarantined as "not a PDF"; every later
    # upload was then skipped as a duplicate with no paper.
    from litforge.ingest import ingest_file
    pid = create_project(conn, paths, "docs")
    doc = _docx(tmp_path / "draft.docx", [("Title", "Neural mass models"), ("", "Text.")])
    wrong = tmp_path / "draft.do"
    wrong.write_bytes(doc.read_bytes())
    assert ingest_file(conn, paths, wrong, pid).status == "quarantined"
    report, results, parsed = add_pdfs(conn, paths, config, pid, [doc], "again")
    assert [r.status for r in results] == ["new"] and results[0].paper_id and parsed[0].status == "ok"


def test_ingest_command_and_removals_handle_documents(conn: sqlite3.Connection, paths: Paths, config: Config,
                                                      tmp_path: Path) -> None:
    from litforge.export.markdown import remove_export
    from litforge.ingest import ingest_paths
    from litforge.library_cleanup import remove_papers
    pid = create_project(conn, paths, "docs")
    doc = _docx(tmp_path / "draft.docx", [("Title", "Neural mass models"), ("", "Text.")])
    [res] = ingest_paths(conn, paths, [doc], pid)  # `litforge ingest draft.docx`
    assert res.status == "new" and res.paper_id
    stored = paths.root / conn.execute("SELECT pdf_path FROM artifacts WHERE id=?", (res.artifact_id,)).fetchone()[0]
    assert stored.suffix == ".docx"
    out, _ = export_markdown(conn, paths, "docs", pid, config=config)
    assert list((out.parent / "pdfs").glob("*.docx"))
    remove_export(out.parent)
    assert not (out.parent / "pdfs").exists()  # review: the copied .docx was left behind
    conn.execute("DELETE FROM project_papers WHERE paper_id=?", (res.paper_id,))
    remove_papers(conn, paths, [res.paper_id])
    assert not stored.exists()  # review: removing the paper left the .docx in the library
